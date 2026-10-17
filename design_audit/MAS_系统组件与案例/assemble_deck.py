"""Assemble image-tool diagrams with user-requested original screenshots and exact data.

This adapts codex-ppt's assembly/notes conventions. Native pages are never labelled
as image-model outputs. Screenshot blobs are embedded unchanged; crop is PPT metadata.
"""
from pathlib import Path
import argparse
import hashlib
import importlib.util
import json
import re
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT=Path(__file__).resolve().parent
OLD=ROOT.parent/'workbench_presentation_2026-10-17'
FONT='Microsoft YaHei'
INK='082646'; BLUE='2086D5'; LIGHT='E6F3FD'; MUTED='5E7489'; LINE='C9DEEE'; WHITE='FFFFFF'
ASSETS={
 'settings':OLD/'presentation_source/assets/24-formal-settings-mid.jpg',
 'task':ROOT/'assets/task-setup-current.jpg',
 'collection':ROOT/'assets/evidence-selection-current.jpg',
 'criteria':ROOT/'assets/criteria-current.jpg',
 'score':OLD/'screenshots/score-desktop.png',
 'leaderboard':OLD/'presentation_source/assets/leaderboard.jpg',
 'evidence':OLD/'screenshots/evidence-fgfr1-grouped.png',
 'export':ROOT/'assets/results-export-current.jpg',
 'upload':OLD/'presentation_source/assets/csv-import.jpg',
}
SCORE=json.loads((OLD/'presentation_source/data/score-details.json').read_text('utf-8-sig'))
SPEC=json.loads((ROOT/'deck_spec.json').read_text('utf-8'))
SPEECH=(ROOT/'speech.md').read_text('utf-8')
BLOCKS={int(m.group(1)):m.group(2).strip() for m in re.finditer(r'^## Slide (\d+):[^\n]*\n(.*?)(?=^## Slide |\Z)',SPEECH,re.M|re.S)}
asset_uses=[]

def rgb(v): return RGBColor.from_string(v)
def east_asian_font(props):
    for tag in ('a:ea','a:cs'):
        el=props.find(qn(tag))
        if el is None:el=OxmlElement(tag);props.append(el)
        el.set('typeface',FONT)
def text(slide, value, x,y,w,h, size=22,bold=False,color=INK,align=PP_ALIGN.LEFT):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=0; tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(str(value).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.alignment=align; p.space_before=Pt(0); p.space_after=Pt(0)
        p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
        east_asian_font(p.font._element)
    return sh

def rect(slide,x,y,w,h,fill=LIGHT,border=None):
    sh=slide.shapes.add_shape(1,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
    if border: sh.line.color.rgb=rgb(border); sh.line.width=Pt(1)
    else: sh.line.fill.background()
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    return sh

def heading(slide,title):
    # Match the approved sample's navy title and generous white margin.
    size=36 if len(title)<25 else 32
    text(slide,title,.42,.35,12.45,.73,size,True)

def picture(slide,key,x,y,w,h,crop=(0,0,1,1)):
    p=ASSETS[key]; iw,ih=Image.open(p).size
    l,t,r,b=crop; ratio=iw*(r-l)/(ih*(b-t))
    rw=min(w,h*ratio);rh=rw/ratio
    x+=(w-rw)/2;y+=(h-rh)/2
    pic=slide.shapes.add_picture(str(p),Inches(x),Inches(y),width=Inches(rw),height=Inches(rh))
    pic.crop_left=l;pic.crop_top=t;pic.crop_right=1-r;pic.crop_bottom=1-b
    asset_uses.append({'slide':len(prs.slides),'key':key,'path':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'crop':[l,t,r,b]})
    return pic

def foot(slide,label): text(slide,label,.5,7.02,12.3,.24,11,color=MUTED)
def note(slide,head,body,y):
    text(slide,head,10.35,y,2.4,.45,24,True,BLUE)
    text(slide,body,10.35,y+.5,2.4,1,17,color=INK)

def line_chart(slide,categories,values,x,y,w,h):
    data=CategoryChartData(); data.categories=categories; data.add_series('工作项',values)
    chart=slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS,Inches(x),Inches(y),Inches(w),Inches(h),data).chart
    chart.has_legend=False; chart.has_title=False
    chart.font.name=FONT;chart.font.size=Pt(17)
    va=chart.value_axis;va.minimum_scale=0;va.maximum_scale=2100;va.major_unit=500
    va.tick_labels.font.size=Pt(13);va.tick_labels.font.color.rgb=rgb(MUTED)
    va.has_major_gridlines=True;va.major_gridlines.format.line.color.rgb=rgb(LINE)
    va.format.line.fill.background()
    ca=chart.category_axis;ca.tick_labels.font.size=Pt(17);ca.tick_labels.font.color.rgb=rgb(INK)
    ca.format.line.color.rgb=rgb(LINE)
    plot=chart.plots[0];plot.has_data_labels=True
    plot.data_labels.position=XL_LABEL_POSITION.ABOVE
    plot.data_labels.font.name=FONT;plot.data_labels.font.size=Pt(22);plot.data_labels.font.bold=True
    plot.data_labels.font.color.rgb=rgb(INK)
    s=chart.series[0];s.format.line.color.rgb=rgb(BLUE);s.format.line.width=Pt(3)
    s.marker.size=9;s.marker.format.fill.solid();s.marker.format.fill.fore_color.rgb=rgb(BLUE)
    s.marker.format.line.color.rgb=rgb(BLUE)
    for props in chart._chartSpace.xpath('.//a:defRPr | .//a:rPr | .//a:endParaRPr'):
        east_asian_font(props)
    return chart

def contribution_bar(slide):
    rows=SCORE['contributions']; groups=[('五靶点','measured','2086D5'),('结构计算','derived','64ADE2'),('模型预测','predicted','B8DAF2')]
    x=.6; y=6.12; total=sum(r['final_score_contribution'] or 0 for r in rows)
    for title,basis,color in groups:
        val=sum(r['final_score_contribution'] or 0 for r in rows if r['determination_basis']==basis)
        width=12.1*val/total
        rect(slide,x,y,width,.32,color)
        text(slide,f'{title}  {val:.5f}',x,y+.42,width,.35,17,True,INK,PP_ALIGN.CENTER)
        x+=width

def native(slide,n):
    if n==16:
        text(slide,'文献与推理：mid  ·  4 并发  ·  本次上限 $5.5',.5,1.27,12.1,.5,25,True,BLUE)
        picture(slide,'settings',.5,1.95,12.25,4.55,(.216,.50,.967,.854))
        foot(slide,'历史运行中的 Settings 画面；最终估计模型费用约 $4.2101。')
    elif n==17:
        picture(slide,'task',.45,1.35,9.35,5.45,(.31,.22,.975,.995))
        note(slide,'研究目标','单一口服药物\n覆盖五个靶点',1.55)
        note(slide,'研究范围','CSF1R、VEGFR1–3\nFGFR1',3.25)
        note(slide,'硬条件','明确填写口服要求',5.05)
        foot(slide,'重新打开已保存配置的填写演示；新研究从空标准集建立。')
    elif n==18:
        picture(slide,'collection',.45,1.35,9.25,5.45,(.31,.15,.972,.91))
        note(slide,'3,177','候选化合物',1.65)
        note(slide,'164 / 3,157','来源 / 保存事实',3.4)
        note(slide,'2026-10-15','输入集合截止时间',5.15)
        foot(slide,'仅使用已有证据。生成阶段去除 136 条明确提及参考候选的事实，剩余 3,021 条。')
    elif n==19:
        picture(slide,'settings',.45,1.35,12.35,4.7,(.215,.015,.969,.474))
        text(slide,'2 / 8 阶段完成',.65,6.25,3.75,.55,25,True,BLUE)
        text(slide,'17 / 18 当前批次',4.7,6.25,4.4,.55,25,True,BLUE)
        text(slide,'时间、活动、预算',9.15,6.25,3.7,.55,23,True,BLUE)
        foot(slide,'历史运行截图：94% 表示当前维度的证据批次进度。')
    elif n==20:
        line_chart(slide,['起点','任务要求','靶点覆盖','口服性质','代谢稳定性','安全性'],[0,6,293,504,1192,1917],.55,1.52,12.1,4.9)
        text(slide,'6 项任务要求 + 1,911 条提议',.65,6.52,11.8,.4,23,True,BLUE)
        foot(slide,'统一口径：已注册标准与未注册提议的工作项总数。')
    elif n==21:
        line_chart(slide,['展开完成','第 1 轮','第 2 轮','第 3 轮'],[1917,147,68,53],.55,1.5,12.1,4.75)
        text(slide,'检查支持  →  明确测量  →  回应质疑  →  合并重复',.65,6.42,12,.43,23,True,BLUE)
        foot(slide,'三轮末端记录：147 → 68 → 53；变化理由与原提议保存在标准历史中。')
    elif n==23:
        picture(slide,'criteria',.45,1.35,12.35,4.85,(.217,.55,.97,.982))
        text(slide,'18 核心',.75,6.27,3.7,.5,28,True,BLUE)
        text(slide,'13 参考',4.95,6.27,3.7,.5,28,True,BLUE)
        text(slide,'22 未启用',9.15,6.27,3.7,.5,28,True,MUTED)
        foot(slide,'31 条进入评估 = 18 核心 + 13 参考；参考项没有筛选或排名权重。')
    elif n==25:
        picture(slide,'score',.5,1.35,12.3,2.45,(.39,.08,.97,.238))
        text(slide,'口服给药有直接记录',.65,4.1,11.9,.52,27,True,BLUE)
        text(slide,'“In both cohorts, surufatinib was administered orally\nat 250 mg once daily.”',.65,4.8,11.9,1.0,24,color=INK)
        text(slide,'五靶点直接证据齐全，满足本次任务要求',.65,6.18,12,.5,25,True)
        foot(slide,'给药记录来源：10.1080/07853890.2026.2698360，sec2.3.p1；本任务未设统一效力硬阈值。')
    elif n==27:
        picture(slide,'score',.45,1.36,12.35,4.5,(.389,.278,.975,.515))
        contribution_bar(slide)
        foot(slide,'15 项有可用值，权重均为 1；直接偏好覆盖为 10/17，预测不计入其分子。')
    elif n==28:
        picture(slide,'leaderboard',.45,1.35,12.35,5.45,(.212,.26,.972,.995))
        foot(slide,'合格榜局部。完整库 3,177：合格 10，待补证据 3,167，淘汰 0，未评估 0。')
    elif n==29:
        picture(slide,'evidence',.5,1.36,12.25,2.3,(.218,.234,.943,.465))
        picture(slide,'evidence',.5,3.8,12.25,2.1,(.233,.55,.932,.727))
        text(slide,'34 nM → 7.4685',.65,6.15,5.3,.5,25,True,BLUE)
        text(slide,'四条观测中位数 → 7.6443',6.15,6.15,6.5,.5,25,True,BLUE)
        foot(slide,'FGFR1 观测：7.2800、7.4685、7.82、7.82；原文 DOI 10.1038/s42004-023-01084-0。')
    elif n==30:
        picture(slide,'export',.45,1.35,12.35,4.4,(.212,.31,.973,.74))
        text(slide,'CSV：完整 3,177 行',.65,6.0,4.35,.6,24,True,BLUE)
        text(slide,'JSON：完整结果',5.05,6.0,3.95,.6,24,True,BLUE)
        text(slide,'历史：标准与来源',9.02,6.0,3.95,.6,24,True,BLUE)
        foot(slide,'按钮导出 CSV；完整 JSON、标准演化及来源记录随本次结果一起交付。')
    elif n==32:
        picture(slide,'upload',.45,1.35,9.5,5.5,(.315,.28,.973,.995))
        note(slide,'候选名单','定义这次比较谁',1.65)
        note(slide,'测量与证据','身份、性质、数值\n单位、条件、来源',3.35)
        note(slide,'新研究','重新形成标准\n再筛选与排序',5.25)
        foot(slide,'药物、载体或其他对象需要匹配相应数据与工具；载体为应用示意。')
    elif n in (33,34): score_table(slide,n)

def score_table(slide,n):
    rows=SCORE['contributions']
    wanted=['CSF1R','VEGFR1','VEGFR2','VEGFR3','FGFR1','Molecular weight','Predicted logP','Oral bioavailability','Time to maximum']
    selected=[]
    for key in wanted:
        matches=[r for r in rows if key.lower() in r['name'].lower()]
        if matches: selected.append(matches[0])
    if n==34:selected=[r for r in rows if r not in selected]
    assert len(selected)==(9 if n==33 else 8), [(r['name']) for r in selected]
    names={
      'CSF1R curated potency':'CSF1R','FGFR1 curated potency':'FGFR1','VEGFR1 curated potency':'VEGFR1','VEGFR2 curated potency':'VEGFR2','VEGFR3 curated potency':'VEGFR3',
      'Molecular weight (Da)':'分子量 ≤ 500 Da','Predicted logP (RDKit)':'logP（RDKit）',
      'Predicted BRENK alert count':'BRENK 警示计数','Predicted NIH alert count':'NIH 警示计数','Predicted PAINS alert count':'PAINS 警示计数',
      'Predicted blood-brain barrier penetration probability (BBB_Martins)':'血脑屏障穿透概率',
      'Predicted CYP1A2 inhibition probability':'CYP1A2 抑制概率','Predicted CYP2C19 inhibition probability':'CYP2C19 抑制概率',
      'Predicted CYP2C9 inhibition probability':'CYP2C9 抑制概率','Predicted CYP2C9 substrate probability':'CYP2C9 底物概率'}
    table=slide.shapes.add_table(len(selected)+1,5,Inches(.5),Inches(1.45),Inches(12.3),Inches(4.95)).table
    for col,width in zip(table.columns,[4.6,2.1,1.65,1.8,2.15]): col.width=Inches(width)
    headers=['偏好','采用值','依据','比较分','总分贡献']
    for j,v in enumerate(headers):table.cell(0,j).text=v
    for i,r in enumerate(selected,1):
        name=names.get(r['name'],r['name'])
        if 'bioavailability' in name.lower():name='人类实测口服 F%'
        if 'maximum plasma' in name.lower():name='人类实测 Tmax'
        value='缺失' if r['assessed_value'] is None else f"{r['assessed_value']:.4f}".rstrip('0').rstrip('.')
        basis={'measured':'实测','derived':'计算','predicted':'预测','absent':'缺失'}.get(r['determination_basis'],r['determination_basis'])
        d='—' if r['desirability'] is None else f"{r['desirability']:.4f}"
        c='—' if r['final_score_contribution'] is None else f"{r['final_score_contribution']:.5f}"
        for j,v in enumerate([name,value,basis,d,c]):table.cell(i,j).text=v
    for i,row in enumerate(table.rows):
        for j,cell in enumerate(row.cells):
            cell.fill.solid();cell.fill.fore_color.rgb=rgb(LIGHT if i==0 else WHITE if i%2 else 'F5F9FC')
            cell.margin_left=Inches(.12);cell.margin_right=Inches(.10);cell.margin_top=Inches(.07);cell.margin_bottom=Inches(.07)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.font.name=FONT;p.font.size=Pt(17 if i else 18);p.font.bold=i==0;p.font.color.rgb=rgb(INK)
                east_asian_font(p.font._element)
                p.alignment=PP_ALIGN.LEFT if j==0 else PP_ALIGN.CENTER
    label='靶点单位 pChEMBL；分子量单位 Da。每项权重 1；15 项平均，直接覆盖 10/17。' if n==33 else '每项权重 1；15 项可用值进入得分平均。直接证据覆盖另算 10/17。'
    text(slide,label,.65,6.65,12,.5,16 if n==33 else 18,color=MUTED)

parser=argparse.ArgumentParser();parser.add_argument('--draft',action='store_true');args=parser.parse_args()
jobs=json.loads((ROOT/'slide_jobs.json').read_text('utf-8'))
assert len(BLOCKS)==34
if not args.draft:
    entries=jobs.get('slides',[])
    assert all(s['status'] in ('accepted','recorded') for s in entries), 'Incomplete image jobs'
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
for info in SPEC['all_slides']:
    n=info['number'];slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb(WHITE)
    if info['render_mode']=='generated':
        img=ROOT/info['source']
        if img.exists():slide.shapes.add_picture(str(img),0,0,width=prs.slide_width,height=prs.slide_height)
        elif args.draft:heading(slide,info['title']);text(slide,'制作中',.5,2,12,.8,34)
        else:raise FileNotFoundError(img)
    else:heading(slide,info['title']);native(slide,n)
    slide.notes_slide.notes_text_frame.text=BLOCKS[n]
out=ROOT/('MAS_系统组件与案例.draft.pptx' if args.draft else 'MAS_系统组件与案例.pptx')
prs.save(out)
with zipfile.ZipFile(out) as z:
    media_hashes={hashlib.sha256(z.read(k)).hexdigest() for k in z.namelist() if k.startswith('ppt/media/')}
assert all(a['sha256'] in media_hashes for a in asset_uses), 'A source screenshot blob was altered'
(ROOT/'asset_usage.json').write_text(json.dumps(asset_uses,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'path':str(out),'slides':len(prs.slides),'notes':len(BLOCKS),'screenshots_unchanged':True,'bytes':out.stat().st_size},ensure_ascii=False))
