"""Translate the plotted vector artists into native PowerPoint objects."""
from pathlib import Path
import json
import numpy as np
from matplotlib.path import Path as MPath
from matplotlib.text import Text
from matplotlib.lines import Line2D
from matplotlib.patches import Patch
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from matplotlib.colors import to_rgba


def export_pptx(fig, path, caption):
    fig.canvas.draw(); renderer=fig.canvas.get_renderer()
    dpi=fig.dpi; fw,fh=fig.get_size_inches()
    prs=Presentation();prs.slide_width=Inches(fw);prs.slide_height=Inches(fh)
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    def coord(x,y):return Inches(float(x)/dpi),Inches(fh-float(y)/dpi)
    def color(value):
        c=to_rgba(value);return RGBColor(*(round(v*255) for v in c[:3])),c[3]
    def poly(points,closed,fill,edge,lw):
        if len(points)<2:return
        converted=[coord(*point) for point in points]
        b=slide.shapes.build_freeform(*converted[0],scale=1)
        b.add_line_segments(converted[1:],close=closed);shape=b.convert_to_shape()
        fc,fa=color(fill);ec,ea=color(edge)
        if closed and fa:shape.fill.solid();shape.fill.fore_color.rgb=fc
        else:shape.fill.background()
        if ea and lw:shape.line.color.rgb=ec;shape.line.width=Pt(lw)
        else:shape.line.fill.background()
    def path_segments(path,fill,edge,lw):
        points=[]
        for vertices,code in path.iter_segments(curves=False,simplify=False):
            if code==MPath.MOVETO:
                if points:poly(points,False,fill,edge,lw)
                points=[vertices[:2]]
            elif code==MPath.CLOSEPOLY:
                poly(points,True,fill,edge,lw);points=[]
            elif code==MPath.STOP:break
            else:points.append(vertices[-2:])
        if points:poly(points,False,fill,edge,lw)
    for ax in fig.axes:
        artists=list(ax.patches)+list(ax.lines)+list(ax.texts)
        if ax.axison:
            artists += list(ax.get_xgridlines())+list(ax.get_ygridlines())
            artists += list(ax.get_xticklabels())+list(ax.get_yticklabels())
            artists += [ax.xaxis.label,ax.yaxis.label]+list(ax.spines.values())
        for artist in sorted(artists,key=lambda a:a.get_zorder()):
            if not artist.get_visible():continue
            if isinstance(artist,Text):
                value=artist.get_text()
                if not value:continue
                bbox=artist.get_window_extent(renderer);padding=1.5
                # Reserve room for PowerPoint's font metrics while retaining
                # the original left, centre or right text anchor.
                extra=bbox.width*.17
                x=bbox.x0-padding-(extra/2 if artist.get_ha()=='center' else extra if artist.get_ha()=='right' else 0)
                left,top=coord(x,bbox.y1+padding)
                shape=slide.shapes.add_textbox(left,top,Inches((bbox.width+2*padding+extra)/dpi),Inches((bbox.height+2*padding)/dpi))
                tf=shape.text_frame;tf.clear();tf.word_wrap=False;tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
                for n,line in enumerate(value.split('\n')):
                    p=tf.paragraphs[0] if n==0 else tf.add_paragraph();p.text=line
                    p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}.get(artist.get_ha(),PP_ALIGN.LEFT)
                    p.space_before=p.space_after=Pt(0);p.line_spacing=1.05
                    p.font.name='Arial';p.font.size=Pt(artist.get_fontsize());p.font.bold=artist.get_weight() in ('bold','semibold',700)
                    p.font.italic=False;p.font.color.rgb=color(artist.get_color())[0]
            elif isinstance(artist,Line2D):
                if artist.get_linestyle() in ('None','none','',None):continue
                points=artist.get_transform().transform(artist.get_xydata())
                for a,b in zip(points,points[1:]):
                    if not np.isfinite([*a,*b]).all():continue
                    shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,*coord(*a),*coord(*b))
                    shape.line.color.rgb=color(artist.get_color())[0];shape.line.width=Pt(artist.get_linewidth())
                    if artist.is_dashed():
                        dash=OxmlElement('a:prstDash');dash.set('val','dash');shape.line._get_or_add_ln().append(dash)
            elif isinstance(artist,Patch):
                path_segments(artist.get_path().transformed(artist.get_transform()),artist.get_facecolor(),artist.get_edgecolor(),artist.get_linewidth())
    for shape in slide.shapes:
        for effect in shape._element.xpath('.//a:effectRef'):effect.set('idx','0')
        props=shape._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}spPr')
        if props is not None:props.append(OxmlElement('a:effectLst'))
    slide.notes_slide.notes_text_frame.text=caption
    prs.save(path)
