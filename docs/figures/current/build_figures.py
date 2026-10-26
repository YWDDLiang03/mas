"""Rebuild the accepted layouts with current terminology and candidate groups."""
import argparse
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import zipfile

from PIL import Image
from pptx import Presentation

out=Path(__file__).resolve().parent
parser=argparse.ArgumentParser()
parser.add_argument('--pdf-python',default=sys.executable,help='Python with PyMuPDF installed')
args=parser.parse_args()
for script in ('restore_original_layout.py','draw_results.py'):
    subprocess.run([sys.executable,str(out/script)],check=True)
subprocess.run(['powershell.exe','-NoProfile','-ExecutionPolicy','Bypass','-File',str(out/'validate_powerpoint.ps1')],check=True)
subprocess.run([args.pdf_python,str(out/'export_workflow_vectors.py')],check=True)

# The selection report and saved score data are the figure's numerical sources.
data=json.loads((out/'figure_data.json').read_text(encoding='utf-8'))
counts=data['selection_reporting_v3']
texts={name:'\n'.join(shape.text for shape in Presentation(out/name).slides[0].shapes
                   if shape.has_text_frame)
       for name in ('TRIAGE_workflow_editable.pptx','Surufatinib_results_editable.pptx')}
for content in texts.values():
    for label,key in (('Evidence incomplete','evidence_incomplete'),('Excluded','excluded')):
        assert f'{label} ({counts[key]:,})' in content
    assert 'Eligibility requirements (6)' in content
    assert 'Partial support; no condition failed' in content
    assert 'Undetermined' not in content and 'Mandatory requirement' not in content
assert f'Eligible ({counts["eligible"]})' in texts['TRIAGE_workflow_editable.pptx']
assert counts['eligible']+counts['evidence_incomplete']+counts['excluded']==data['candidate_count']
powerpoint=json.loads((out/'powerpoint-validation.json').read_text(encoding='utf-8-sig'))
assert powerpoint['passed'], 'PowerPoint text overflow'
layout=json.loads((out/'workflow-layout-preservation.json').read_text(encoding='utf-8'))
files={}
for name in texts:
    with zipfile.ZipFile(out/name) as package:
        slides=b'\n'.join(package.read(p) for p in package.namelist()
                          if p.startswith('ppt/slides/slide') and p.endswith('.xml'))
        stats={'embedded_images':sum(p.startswith('ppt/media/') for p in package.namelist()),
               'gradient_fills':slides.count(b'<a:gradFill'),
               'shadows':slides.count(b'<a:outerShdw')+slides.count(b'<a:innerShdw')}
        assert not any(stats.values()), (name,stats)
    files[name]={'sha256':hashlib.sha256((out/name).read_bytes()).hexdigest(),**stats}
report={'original_layouts_preserved':not layout['unplanned_geometry_changes'],
        'workflow_layout':layout,'files':files,'png_dimensions':{
            name:list(Image.open(out/f'{name}.png').size)
            for name in ('figure_1_workflow','figure_2_surufatinib_results')},
        'powerpoint':powerpoint,'selection_counts':counts,
        'eligibility_requirement_semantics':{'condition_requirements':1,'evidence_requirements':5,
            'scoring_criteria':17,'shared_admission_scoring':5,'unique_decision_rules':18}}
(out/'validation.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
(out/'workflow-native.pdf').unlink(missing_ok=True)
print('Figure labels, source counts, editability and PowerPoint text fit verified.')
